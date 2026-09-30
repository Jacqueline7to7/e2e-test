import streamlit as st
import pandas as pd
import io
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border
from copy import copy

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.join(BASE_DIR, "templates")

st.set_page_config(
    page_title="E2E 自动化交付平台",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
    .stApp { background: #F9FAFB; }
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    .main .block-container { max-width: 1400px; padding: 24px 32px; }
    .page-title { font-size: 24px; font-weight: 700; color: #111827; margin: 0 0 4px 0; }
    .page-subtitle { font-size: 13px; color: #6B7280; margin: 0 0 24px 0; }
    .kpi-card {
        background: #FFFFFF; border: 1px solid #E5E7EB; border-radius: 12px;
        padding: 20px; box-shadow: 0 1px 2px rgba(0,0,0,0.05);
    }
    .kpi-label { font-size: 13px; color: #6B7280; margin-bottom: 8px; }
    .kpi-value { font-size: 30px; font-weight: 700; color: #111827; line-height: 1.2; }
    .kpi-trend { font-size: 12px; margin-top: 8px; }
    .e2e-card {
        background: #FFFFFF; border: 1px solid #E5E7EB; border-radius: 12px;
        padding: 20px; box-shadow: 0 1px 2px rgba(0,0,0,0.05); margin-bottom: 16px;
    }
    .check-pass { color: #16A34A; font-size: 14px; padding: 4px 0; }
    .check-fail { color: #DC2626; font-size: 14px; padding: 4px 0; }
    .check-warn { color: #D97706; font-size: 14px; padding: 4px 0; }
</style>
""", unsafe_allow_html=True)

if "uploaded_data" not in st.session_state:
    st.session_state["uploaded_data"] = []
if "fetch_df" not in st.session_state:
    st.session_state["fetch_df"] = None
if "backfill_df" not in st.session_state:
    st.session_state["backfill_df"] = None
if "result_df" not in st.session_state:
    st.session_state["result_df"] = None
if "qc_issues" not in st.session_state:
    st.session_state["qc_issues"] = []

# ============ 校验函数 ============
RED_REQUIRED = [
    "Brand", "Data Period", "Product", "Category", "Sub-category",
    "Campaign Type", "Campaign Period", "跨域项目名称",
    "计划有效周期", "笔记发布日期", "达人名称", "笔记ID",
    "达人Spending", "合约Spending", "竞价Spending",
    "品线天猫商品ID List", "TM单品主链客单价", "TM全店客单价",
]
DY_REQUIRED = [
    "Data Period", "Category", "Sub-category", "Campaign type", "Product",
    "星图订单ID", "达人名称", "发布日期", "KOL花费",
    "CPC Paid Media花费", "CPM Paid Media花费（热推）",
    "竞价投流花费（种草通）", "品线抖音商品ID List",
    "DY单品主链客单价", "DY全店客单价",
]

def validate_dataframe(df, platform):
    results = []
    required = RED_REQUIRED if platform == "RED" else DY_REQUIRED
    missing_cols = [c for c in required if c not in df.columns]
    if missing_cols:
        return "FAIL", [("FAIL", f"缺少必填列：{', '.join(missing_cols)}")]
    results.append(("PASS", f"必填列完整（共 {len(required)} 列）"))

    empty_cols = [c for c in required if df[c].isnull().any() or (df[c].astype(str).str.strip() == "").any()]
    if empty_cols:
        results.append(("FAIL", f"以下列存在空值：{', '.join(empty_cols)}"))
    else:
        results.append(("PASS", "所有必填列无空值"))

    spend_cols = [c for c in df.columns if "Spending" in c or "花费" in c]
    neg_cols = [c for c in spend_cols if (pd.to_numeric(df[c], errors="coerce") < 0).any()]
    if neg_cols:
        results.append(("FAIL", f"以下花费列存在负值：{', '.join(neg_cols)}"))
    else:
        results.append(("PASS", "花费列均为非负"))

    aov_cols = [c for c in df.columns if "客单价" in c]
    bad_aov = [c for c in aov_cols if (pd.to_numeric(df[c], errors="coerce") <= 0).any()]
    if bad_aov:
        results.append(("WARN", f"以下客单价列存在 <=0 的值：{', '.join(bad_aov)}"))
    else:
        results.append(("PASS", "客单价列均 > 0"))

    id_col = "笔记ID" if platform == "RED" else "星图订单ID"
    if id_col in df.columns:
        dup = df[df.duplicated(subset=[id_col], keep=False)]
        if len(dup) > 0:
            results.append(("WARN", f"{id_col} 存在重复：{dup[id_col].nunique()} 个重复值"))
        else:
            results.append(("PASS", f"{id_col} 无重复"))

    results.append(("PASS", f"共 {len(df)} 行数据，{len(df.columns)} 列"))
    if any(r[0] == "FAIL" for r in results):
        return "FAIL", results
    elif any(r[0] == "WARN" for r in results):
        return "WARN", results
    return "PASS", results

# ============ 取数清单生成 ============
def build_fetch_list(uploaded_data):
    rows = []
    idx = 1
    for item in uploaded_data:
        df = item.get("df")
        if df is None:
            continue
        plat = item["平台"]
        for _, r in df.iterrows():
            if plat == "RED":
                brand = r.get("Brand", "")
                product = r.get("Product", "")
                project = r.get("跨域项目名称", product)
                order = r.get("笔记ID", "")
                period = r.get("计划有效周期", "") or r.get("Campaign Period", "")
                read_dim = "红书阅读数"
            else:
                brand = r.get("Brand", r.get("Product", ""))
                product = r.get("Product", "")
                project = product
                order = r.get("星图订单ID", "")
                period = r.get("Data Period", "")
                read_dim = "视频观看数"

            for dim in [read_dim, "进店UV", "购买人数", "渠道新客数"]:
                rows.append({
                    "#": idx,
                    "品牌": str(brand),
                    "平台": plat,
                    "项目名称": str(project),
                    "订单名称": str(order),
                    "时间范围": str(period),
                    "维度": dim,
                    "窗口": "R30",
                    "状态": "待查询",
                    "数值": None,
                })
                idx += 1
    return pd.DataFrame(rows)

# ============ 回填匹配 ============
MATCH_KEYS = ["品牌", "平台", "项目名称", "订单名称", "时间范围", "维度"]

def match_backfill(fetch_df, backfill_df):
    if fetch_df is None or backfill_df is None:
        return None
    missing = [k for k in MATCH_KEYS + ["数值"] if k not in backfill_df.columns]
    if missing:
        return {"error": f"回填文件缺少列：{', '.join(missing)}"}

    merged = fetch_df.merge(
        backfill_df[MATCH_KEYS + ["数值"]],
        on=MATCH_KEYS, how="left", suffixes=("", "_回填"),
    )
    if "数值_回填" in merged.columns:
        merged["数值"] = merged["数值_回填"].combine_first(merged["数值"])
        merged = merged.drop(columns=["数值_回填"])
    merged["状态"] = merged.apply(
        lambda r: "已完成" if pd.notna(r["数值"]) else r["状态"], axis=1
    )
    return merged

# ============ 聚合上传数据 ============
def aggregate_uploaded(uploaded_data):
    rows = []
    for item in uploaded_data:
        df = item.get("df")
        plat = item["平台"]
        if df is None:
            continue
        for _, r in df.iterrows():
            category = str(r.get("Category", ""))
            if plat == "RED":
                brand = r.get("Brand", "")
                product = r.get("Product", "")
                project = r.get("跨域项目名称", product)
                spend = sum([
                    pd.to_numeric(r.get("达人Spending", 0), errors="coerce") or 0,
                    pd.to_numeric(r.get("合约Spending", 0), errors="coerce") or 0,
                    pd.to_numeric(r.get("竞价Spending", 0), errors="coerce") or 0,
                ])
                aov = pd.to_numeric(r.get("TM全店客单价", 0), errors="coerce") or 0
                item_aov = pd.to_numeric(r.get("TM单品主链客单价", 0), errors="coerce") or 0
            else:
                brand = r.get("Brand", r.get("Product", ""))
                product = r.get("Product", "")
                project = product
                spend = sum([
                    pd.to_numeric(r.get("KOL花费", 0), errors="coerce") or 0,
                    pd.to_numeric(r.get("CPC Paid Media花费", 0), errors="coerce") or 0,
                    pd.to_numeric(r.get("CPM Paid Media花费（热推）", 0), errors="coerce") or 0,
                    pd.to_numeric(r.get("竞价投流花费（种草通）", 0), errors="coerce") or 0,
                ])
                aov = pd.to_numeric(r.get("DY全店客单价", 0), errors="coerce") or 0
                item_aov = pd.to_numeric(r.get("DY单品主链客单价", 0), errors="coerce") or 0
            rows.append({
                "品牌": str(brand), "平台": plat, "项目名称": str(project),
                "产品": str(product), "Category": category,
                "媒体花费": spend, "全店客单价": aov, "单品客单价": item_aov,
            })
    if not rows:
        return pd.DataFrame(columns=["品牌", "平台", "项目名称", "产品", "Category", "媒体花费", "全店客单价", "单品客单价"])
    return pd.DataFrame(rows).groupby(["品牌", "平台", "项目名称"], as_index=False).agg({
        "产品": "first", "Category": "first",
        "媒体花费": "sum", "全店客单价": "first", "单品客单价": "first",
    })

# ============ 计算 E2E 指标 ============
def calculate_metrics(uploaded_data, backfill_df):
    agg = aggregate_uploaded(uploaded_data)
    if len(agg) == 0:
        return None, []

    bf = backfill_df.copy()
    bf["数值"] = pd.to_numeric(bf["数值"], errors="coerce").fillna(0)

    if bf["数值"].sum() == 0:
        return None, [("FAIL", "全局", "回填数据全部为空或为 0，无法计算")]

    pivoted = (
        bf.groupby(["品牌", "平台", "项目名称", "维度"])["数值"]
        .sum()
        .unstack(fill_value=0)
        .reset_index()
    )

    if "红书阅读数" in pivoted.columns:
        pivoted["阅读数"] = pivoted["红书阅读数"]
    elif "视频观看数" in pivoted.columns:
        pivoted["阅读数"] = pivoted["视频观看数"]
    else:
        pivoted["阅读数"] = 0

    for col in ["进店UV", "购买人数", "渠道新客数", "阅读数"]:
        if col not in pivoted.columns:
            pivoted[col] = 0

    result = agg.merge(
        pivoted[["品牌", "平台", "项目名称", "阅读数", "进店UV", "购买人数", "渠道新客数"]],
        on=["品牌", "平台", "项目名称"], how="left",
    ).fillna(0)

    # ===== 关键修复：强制所有数值列转成 numeric =====
    numeric_cols = ["媒体花费", "全店客单价", "单品客单价",
                    "阅读数", "进店UV", "购买人数", "渠道新客数"]
    for c in numeric_cols:
        if c in result.columns:
            result[c] = pd.to_numeric(result[c], errors="coerce").fillna(0).astype(float)

    def safe_div(a, b):
        return (a / b).where(b != 0, 0.0)

    result["进店率"] = safe_div(result["进店UV"], result["阅读数"])
    result["购买转化率"] = safe_div(result["购买人数"], result["进店UV"])
    result["渠道新客占比"] = safe_div(result["渠道新客数"], result["购买人数"])
    result["获客成本"] = safe_div(result["媒体花费"], result["购买人数"])
    result["新客成本"] = safe_div(result["媒体花费"], result["渠道新客数"])
    result["预估新客GMV"] = result["全店客单价"] * result["渠道新客数"]
    result["预估整体新客ROI"] = safe_div(result["预估新客GMV"], result["媒体花费"])

    issues = []
    for _, r in result.iterrows():
        key = f"{r['品牌']} / {r['项目名称']}"
        if r["阅读数"] == 0:
            issues.append(("FAIL", key, "阅读数为 0，无法计算进店率"))
        if r["阅读数"] > 0 and r["进店UV"] > r["阅读数"]:
            issues.append(("FAIL", key, f"进店UV（{int(r['进店UV'])}）> 阅读数（{int(r['阅读数'])}）"))
        if r["进店UV"] > 0 and r["购买人数"] > r["进店UV"]:
            issues.append(("FAIL", key, f"购买人数（{int(r['购买人数'])}）> 进店UV（{int(r['进店UV'])}）"))
        if r["购买人数"] > 0 and r["渠道新客数"] > r["购买人数"]:
            issues.append(("FAIL", key, f"渠道新客数（{int(r['渠道新客数'])}）> 购买人数（{int(r['购买人数'])}）"))
        if r["媒体花费"] <= 0:
            issues.append(("WARN", key, "媒体花费为 0 或负"))
        if r["全店客单价"] <= 0:
            issues.append(("WARN", key, "全店客单价为 0 或负"))
    return result, issues

# ============ 侧边栏 ============
with st.sidebar:
    st.markdown("### 📊 E2E 自动化交付平台")
    st.markdown("---")
    page = st.radio(
        "导航",
        ["📊 概览", "📋 信息收集", "📥 取数清单", "📤 数据回填",
         "🧮 计算质检", "📄 报告输出", "📈 历史趋势"],
        label_visibility="collapsed",
    )
    st.markdown("---")
    st.markdown("**全局筛选**")
    period_filter = st.selectbox("周期", ["2026-08", "2026-07", "2026-06"])
    platform_filter = st.selectbox("平台", ["全部", "RED", "DY"])
    brand_filter = st.selectbox("品牌", ["全部", "品牌A", "品牌B", "品牌C", "品牌D", "品牌E"])
    st.markdown("---")
    st.caption("版本 v0.5")

# ============ 页面：概览 ============
if page == "📊 概览":
    st.markdown('<div class="page-title">E2E 交付概览</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    cols = st.columns(4)
    kpis = [
        ("品牌总数", "12", "全部平台", "#111827"),
        ("已完成", "8", "🟢 67%", "#16A34A"),
        ("待处理", "3", "🟡 25%", "#D97706"),
        ("异常", "1", "🔴 8%", "#DC2626"),
    ]
    for col, (label, value, trend, color) in zip(cols, kpis):
        with col:
            st.markdown(f"""
            <div class="kpi-card">
                <div class="kpi-label">{label}</div>
                <div class="kpi-value" style="color: {color};">{value}</div>
                <div class="kpi-trend">{trend}</div>
            </div>
            """, unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### 各品牌交付状态")
    st.info("待接入真实数据后，这里显示各品牌的交付状态表。")
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### DDL 提醒")
    st.warning("品牌C DMP 取数未完成，剩余 3 项，DDL：09-05 18:00")
    st.error("品牌E 数据异常：新客数 > 购买人数，需复核")
    st.markdown('</div>', unsafe_allow_html=True)

# ============ 页面：信息收集 ============
elif page == "📋 信息收集":
    st.markdown('<div class="page-title">信息收集与校验</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### 下载示例模板")
    st.caption("如果你还没有收集表，可以先下载示例模板，填完后上传测试。")

    col1, col2 = st.columns([1, 3])
    with col1:
        template_type = st.selectbox("模板类型", ["红书", "抖音"])
    with col2:
        st.write("")

    if st.button("生成示例模板"):
        if template_type == "红书":
            sample = pd.DataFrame({
                "Brand": ["品牌A"],
                "Data Period": ["2026.07.01-08.31"],
                "Product": ["产品A"],
                "Category": ["Make up"],
                "Sub-category": ["Lip"],
                "Campaign Type": ["RP+CVD"],
                "Campaign Period": ["7.1-8.19"],
                "跨域项目名称": ["产品A"],
                "计划有效周期": ["7.1-8.19"],
                "笔记发布日期": ["2026-07-15"],
                "达人名称": ["达人小王"],
                "笔记ID": ["note_001"],
                "达人Spending": [1830860],
                "合约Spending": [519374],
                "竞价Spending": [0],
                "品线天猫商品ID List": ["123456789"],
                "TM单品主链客单价": [380],
                "TM全店客单价": [1015],
            })
        else:
            sample = pd.DataFrame({
                "Data Period": ["2026.08"],
                "Category": ["Skincare"],
                "Sub-category": ["Lotion"],
                "Campaign type": ["Repush"],
                "Product": ["产品O"],
                "星图订单ID": ["order_001"],
                "达人名称": ["达人小李"],
                "发布日期": ["2026-08-15"],
                "KOL花费": [200000],
                "CPC Paid Media花费": [50000],
                "CPM Paid Media花费（热推）": [30000],
                "竞价投流花费（种草通）": [10000],
                "品线抖音商品ID List": ["987654321"],
                "DY单品主链客单价": [152],
                "DY全店客单价": [991],
            })

        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            sample.to_excel(writer, index=False, sheet_name="Sheet1")
        buffer.seek(0)

        st.download_button(
            label="⬇ 下载示例模板 Excel",
            data=buffer,
            file_name=f"{template_type}_示例模板.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### 上传收集表")
    uploaded_files = st.file_uploader(
        "上传红书/抖音信息收集表（支持多文件）",
        type=["xlsx", "csv"],
        accept_multiple_files=True,
    )

    if uploaded_files:
        all_results = []
        for f in uploaded_files:
            try:
                if f.name.endswith(".csv"):
                    df = pd.read_csv(f)
                else:
                    df = pd.read_excel(f)

                if "红书" in f.name or "RED" in f.name.upper():
                    plat = "RED"
                elif "抖音" in f.name or "DY" in f.name.upper():
                    plat = "DY"
                else:
                    plat = "RED" if "笔记ID" in df.columns else "DY"

                status, details = validate_dataframe(df, plat)
                all_results.append({
                    "文件名": f.name, "平台": plat, "行数": len(df),
                    "状态": status, "明细": details, "df": df,
                })
            except Exception as e:
                all_results.append({
                    "文件名": f.name, "平台": "未知", "行数": 0,
                    "状态": "FAIL", "明细": [("FAIL", f"文件解析失败：{str(e)}")],
                    "df": None,
                })

        st.session_state["uploaded_data"] = all_results

        total = len(all_results)
        passed = sum(1 for r in all_results if r["状态"] == "PASS")
        warned = sum(1 for r in all_results if r["状态"] == "WARN")
        failed = sum(1 for r in all_results if r["状态"] == "FAIL")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("文件总数", total)
        c2.metric("通过", passed)
        c3.metric("警告", warned)
        c4.metric("失败", failed)

        st.markdown("---")
        st.markdown("#### 校验结果")
        for r in all_results:
            status_icon = {"PASS": "🟢", "WARN": "🟡", "FAIL": "🔴"}[r["状态"]]
            with st.expander(f"{status_icon} {r['文件名']}（{r['平台']}，{r['行数']} 行）"):
                for level, msg in r["明细"]:
                    if level == "PASS":
                        st.markdown(f'<div class="check-pass">✅ {msg}</div>', unsafe_allow_html=True)
                    elif level == "WARN":
                        st.markdown(f'<div class="check-warn">⚠ {msg}</div>', unsafe_allow_html=True)
                    else:
                        st.markdown(f'<div class="check-fail">❌ {msg}</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)

# ============ 页面：取数清单 ============
elif page == "📥 取数清单":
    st.markdown('<div class="page-title">DMP / 云图取数清单</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    if not st.session_state["uploaded_data"]:
        st.warning("请先到「📋 信息收集」页上传收集表，系统会自动生成取数清单。")
    else:
        st.info("💡 请按以下清单到 DMP 后台逐项查询，查询完成后到「📤 数据回填」页上传或粘贴结果。")

        fetch_df = build_fetch_list(st.session_state["uploaded_data"])

        if len(fetch_df) == 0:
            st.warning("未能从收集表中提取到有效数据，请检查收集表格式。")
        else:
            total = len(fetch_df)
            done = len(fetch_df[fetch_df["状态"] == "已完成"])
            pending = len(fetch_df[fetch_df["状态"] == "待查询"])
            error = len(fetch_df[fetch_df["状态"] == "异常"])

            c1, c2, c3, c4 = st.columns(4)
            c1.metric("任务总数", total)
            c2.metric("已完成", done)
            c3.metric("待查询", pending)
            c4.metric("异常", error)

            st.progress(done / total if total > 0 else 0, text=f"取数进度：{done} / {total}")

            st.markdown("#### 取数任务清单")
            st.caption("每个活动对应 4 个维度：阅读数 / 进店UV / 购买人数 / 渠道新客数。")

            edited_df = st.data_editor(
                fetch_df,
                column_config={
                    "状态": st.column_config.SelectboxColumn(
                        "状态",
                        options=["待查询", "已完成", "异常"],
                        required=True, default="待查询",
                    ),
                },
                use_container_width=True, hide_index=True, key="fetch_editor",
            )
            st.session_state["fetch_df"] = edited_df

            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                edited_df.to_excel(writer, index=False, sheet_name="取数清单")
            buffer.seek(0)

            st.download_button(
                label="⬇ 导出取数清单 Excel",
                data=buffer,
                file_name=f"取数清单_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# ============ 页面：数据回填 ============
elif page == "📤 数据回填":
    st.markdown('<div class="page-title">数据回填</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    if st.session_state["fetch_df"] is None:
        st.warning("请先到「📥 取数清单」页生成取数清单。")
    else:
        fetch_df = st.session_state["fetch_df"]

        st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
        st.markdown("#### 下载回填模板")
        st.caption("下载模板后，把 DMP 查询结果填入「数值」列，再上传回来。")

        template = fetch_df[MATCH_KEYS].copy()
        template["数值"] = None

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            template.to_excel(writer, index=False, sheet_name="回填模板")
        buf.seek(0)

        st.download_button(
            label="⬇ 下载回填模板 Excel",
            data=buf,
            file_name=f"回填模板_{datetime.now().strftime('%Y%m%d')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        st.markdown('</div>', unsafe_allow_html=True)

        tab1, tab2 = st.tabs(["上传文件", "手动粘贴"])

        with tab1:
            backfill_file = st.file_uploader("上传回填文件（Excel / CSV）", type=["xlsx", "csv"], key="backfill_upload")
            if backfill_file:
                try:
                    if backfill_file.name.endswith(".csv"):
                        bdf = pd.read_csv(backfill_file)
                    else:
                        bdf = pd.read_excel(backfill_file)
                    result = match_backfill(fetch_df, bdf)
                    if isinstance(result, dict) and "error" in result:
                        st.error(result["error"])
                    else:
                        st.session_state["backfill_df"] = result
                        st.success(f"回填文件解析成功，共 {len(bdf)} 行数据。")
                except Exception as e:
                    st.error(f"文件解析失败：{str(e)}")

        with tab2:
            pasted = st.text_area(
                "粘贴数据（第一行为表头，列用 Tab 分隔）",
                height=200,
                placeholder="品牌\t平台\t项目名称\t订单名称\t时间范围\t维度\t数值\n品牌A\tRED\t产品A\tnote_001\t7.1-8.19\t进店UV\t544798",
            )
            if pasted:
                try:
                    from io import StringIO
                    bdf = pd.read_csv(StringIO(pasted), sep="\t")
                    result = match_backfill(fetch_df, bdf)
                    if isinstance(result, dict) and "error" in result:
                        st.error(result["error"])
                    else:
                        st.session_state["backfill_df"] = result
                        st.success(f"粘贴数据解析成功，共 {len(bdf)} 行数据。")
                except Exception as e:
                    st.error(f"粘贴数据解析失败：{str(e)}")

        if st.session_state["backfill_df"] is not None:
            st.markdown("---")
            st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
            st.markdown("#### 匹配结果")
            result_df = st.session_state["backfill_df"]
            matched = result_df["数值"].notna().sum()
            total = len(result_df)

            c1, c2, c3 = st.columns(3)
            c1.metric("任务总数", total)
            c2.metric("已匹配", matched)
            c3.metric("未匹配", total - matched)

            display_df = result_df.copy()
            display_df["匹配状态"] = display_df["数值"].apply(lambda v: "✅ 已匹配" if pd.notna(v) else "❌ 未匹配")
            st.dataframe(
                display_df[["品牌", "平台", "项目名称", "订单名称", "时间范围", "维度", "数值", "匹配状态"]],
                use_container_width=True, hide_index=True,
            )

            unmatched = display_df[display_df["匹配状态"] == "❌ 未匹配"]
            if len(unmatched) > 0:
                st.warning(f"有 {len(unmatched)} 项未匹配到数据。")

            buf2 = io.BytesIO()
            with pd.ExcelWriter(buf2, engine="openpyxl") as writer:
                display_df.to_excel(writer, index=False, sheet_name="匹配结果")
            buf2.seek(0)
            st.download_button(
                label="⬇ 导出匹配结果 Excel", data=buf2,
                file_name=f"匹配结果_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            st.markdown('</div>', unsafe_allow_html=True)

# ============ 页面：计算质检 ============
elif page == "🧮 计算质检":
    st.markdown('<div class="page-title">计算与质检</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    if st.session_state["uploaded_data"] is None or len(st.session_state["uploaded_data"]) == 0:
        st.warning("请先到「📋 信息收集」页上传收集表。")
    elif st.session_state["backfill_df"] is None:
        st.warning("请先到「📤 数据回填」页完成数据回填。")
    else:
        col1, col2, col3 = st.columns([1, 1, 4])
        with col1:
            if st.button("开始计算", type="primary"):
                with st.spinner("正在计算 E2E 指标..."):
                    result_df, issues = calculate_metrics(
                        st.session_state["uploaded_data"],
                        st.session_state["backfill_df"],
                    )
                    st.session_state["result_df"] = result_df
                    st.session_state["qc_issues"] = issues
                st.success("计算完成！")
        with col2:
            if st.button("重新计算"):
                st.session_state["result_df"] = None
                st.session_state["qc_issues"] = []
                st.info("已清除计算结果，请点击「开始计算」。")

        if st.session_state["result_df"] is not None:
            result_df = st.session_state["result_df"]
            issues = st.session_state["qc_issues"]

            # 质检汇总
            fail_count = sum(1 for i in issues if i[0] == "FAIL")
            warn_count = sum(1 for i in issues if i[0] == "WARN")

            st.markdown("---")
            st.markdown("#### 质检结果")
            c1, c2, c3 = st.columns(3)
            c1.metric("活动总数", len(result_df))
            c2.metric("失败项", fail_count, delta_color="inverse" if fail_count > 0 else "normal")
            c3.metric("警告项", warn_count)

            if fail_count > 0:
                st.error(f"存在 {fail_count} 项失败，需复核后再生成报告。")
            elif warn_count > 0:
                st.warning(f"存在 {warn_count} 项警告，可继续但建议复核。")
            else:
                st.success("全部通过，可以进入报告输出。")

            if issues:
                st.markdown("##### 异常明细")
                for level, key, msg in issues:
                    if level == "FAIL":
                        st.markdown(f'<div class="check-fail">❌ {key}：{msg}</div>', unsafe_allow_html=True)
                    else:
                        st.markdown(f'<div class="check-warn">⚠ {key}：{msg}</div>', unsafe_allow_html=True)

            # 关键指标预览
            st.markdown("---")
            st.markdown("#### 关键指标预览")

            display_cols = ["品牌", "平台", "项目名称", "媒体花费", "阅读数", "进店UV", "购买人数",
                            "渠道新客数", "进店率", "购买转化率", "新客成本", "预估整体新客ROI"]
            display_df = result_df[display_cols].copy()

            # 格式化
            display_df["媒体花费"] = display_df["媒体花费"].apply(lambda x: f"{x:,.0f}")
            display_df["进店率"] = display_df["进店率"].apply(lambda x: f"{x:.2%}")
            display_df["购买转化率"] = display_df["购买转化率"].apply(lambda x: f"{x:.2%}")
            display_df["新客成本"] = display_df["新客成本"].apply(lambda x: f"¥{x:,.0f}")
            display_df["预估整体新客ROI"] = display_df["预估整体新客ROI"].apply(lambda x: f"{x:.2f}")

            st.dataframe(display_df, use_container_width=True, hide_index=True)

            # 导出
            buf = io.BytesIO()
            with pd.ExcelWriter(buf, engine="openpyxl") as writer:
                result_df.to_excel(writer, index=False, sheet_name="E2E指标")
            buf.seek(0)
            st.download_button(
                label="⬇ 导出计算结果 Excel",
                data=buf,
                file_name=f"E2E指标_{datetime.now().strftime('%Y%m%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# ============ 页面：报告输出 ============
elif page == "📄 报告输出":
    import re
    from openpyxl.utils import column_index_from_string, get_column_letter

    st.markdown('<div class="page-title">报告输出</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)

    if st.session_state["result_df"] is None:
        st.warning("请先到「🧮 计算质检」页完成计算。")
    else:
        st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
        st.markdown("#### 生成报告")

        col1, col2 = st.columns(2)
        with col1:
            platform_report = st.selectbox("平台", ["RED", "DY"])
        with col2:
            period_label = st.text_input("周期标签", value="2026.09.30")

        st.caption(f"将在原数据块上方插入 {period_label} 数据块，样式完全复制原块。")

        if st.button("生成报告", type="primary"):
            try:
                result_df = st.session_state["result_df"]

                if platform_report == "RED":
                    df_to_write = result_df[result_df["平台"] == "RED"].copy()
                    template_path = os.path.join(TEMPLATE_DIR, "RED_template.xlsx")
                else:
                    df_to_write = result_df[result_df["平台"] == "DY"].copy()
                    template_path = os.path.join(TEMPLATE_DIR, "DY_template.xlsx")

                if len(df_to_write) == 0:
                    st.error(f"没有找到 {platform_report} 平台的数据。")
                else:
                    wb = openpyxl.load_workbook(template_path)
                    ws = wb["Summary"]

                    # ===== 步骤 1：找到原数据块的行位置 =====
                    old_title_row = None
                    for row in range(1, ws.max_row + 1):
                        val = ws.cell(row=row, column=2).value
                        if val and "Campaign起始日期" in str(val):
                            old_title_row = row
                            break

                    if old_title_row is None:
                        st.error("找不到原 Campaign 标题行，请检查模板结构。")
                        st.stop()

                    old_header_row = None
                    for row in range(old_title_row + 1, ws.max_row + 1):
                        val = ws.cell(row=row, column=2).value
                        if val == "Brand":
                            old_header_row = row
                            break

                    old_data_row = old_header_row + 1

                    # ===== 步骤 2：捕获原块样式 =====
                    NCOLS = 30

                    def capture_style(row):
                        result = {}
                        for c in range(1, NCOLS + 1):
                            cell = ws.cell(row=row, column=c)
                            result[c] = {
                                "font": copy(cell.font),
                                "fill": copy(cell.fill),
                                "border": copy(cell.border),
                                "alignment": copy(cell.alignment),
                                "number_format": cell.number_format,
                            }
                        return result

                    def apply_style(row, style):
                        for c, s in style.items():
                            cell = ws.cell(row=row, column=c)
                            cell.font = copy(s["font"])
                            cell.fill = copy(s["fill"])
                            cell.border = copy(s["border"])
                            cell.alignment = copy(s["alignment"])
                            cell.number_format = s["number_format"]

                    title_style = capture_style(old_title_row)
                    header_style = capture_style(old_header_row)
                    data_style = capture_style(old_data_row)

                    title_height = ws.row_dimensions[old_title_row].height
                    header_height = ws.row_dimensions[old_header_row].height
                    data_height = ws.row_dimensions[old_data_row].height

                    # ===== 步骤 3：保存所有公式 =====
                    formula_map = {}
                    for row in ws.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and cell.value.startswith("="):
                                formula_map[(cell.row, cell.column)] = cell.value

                    # ===== 步骤 4：插入行 =====
                    n_data = len(df_to_write)
                    n_rows = 2 + n_data   # 标题 + 表头 + N 数据
                    insert_row = old_title_row
                    ws.insert_rows(insert_row, n_rows)

                    # ===== 步骤 5：修正所有公式的行号引用 =====
                    pattern = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')

                    def shift_ref(match):
                        cd, col, rd, rn = match.groups()
                        rn = int(rn)
                        if rn >= insert_row:
                            rn += n_rows
                        return f"{cd}{col}{rd}{rn}"

                    for (old_r, old_c), formula in formula_map.items():
                        new_formula = pattern.sub(shift_ref, formula)
                        if old_r >= insert_row:
                            new_r = old_r + n_rows
                            ws.cell(row=new_r, column=old_c).value = new_formula
                        else:
                            ws.cell(row=old_r, column=old_c).value = new_formula

                    # ===== 步骤 6：写新标题行（样式复制原标题行） =====
                    title_row = insert_row
                    apply_style(title_row, title_style)
                    ws.cell(row=title_row, column=2).value = f"Campaign起始日期-{period_label}"
                    if title_height:
                        ws.row_dimensions[title_row].height = title_height

                    # ===== 步骤 7：写新表头行（样式复制原表头行） =====
                    header_row = insert_row + 1
                    apply_style(header_row, header_style)

                    if platform_report == "RED":
                        header_map = {
                            2: "Brand", 3: "Media Format", 4: "媒体花费\n包含KOL+投流",
                            5: "红书阅读数", 6: "归因逻辑", 7: "进店UV\n（=商品页浏览）",
                            8: "进店UV vs 目标", 9: "进店率\n（=进店UV/阅读数)",
                            10: "购买转化率\n(=购买人数/进店数)", 11: "购买人数",
                            12: "渠道新客数", 13: "渠道新客数 vs 目标",
                            14: "渠道新客占比\n(=渠道新客/购买人数)",
                            15: "新客成本\n(=媒体花费/渠道新客数)",
                            16: "预估新客GMV\n(=客单价*新客数)",
                            17: "预估整体新客ROI\n(=新客GMV/媒体花费)",
                            18: "全店客单价",
                        }
                    else:
                        header_map = {
                            2: "Brand", 3: "Media Format", 4: "媒体花费\n包含KOL+投流",
                            5: "视频观看数", 6: "归因逻辑", 7: "进店UV\n（=商品页浏览）",
                            8: "进店UV vs 目标", 9: "进店率\n（=进店UV/阅读数)",
                            10: "购买转化率\n(=购买人数/进店数)", 11: "购买人数",
                            12: "渠道新客数", 13: "渠道新客数 vs 目标",
                            14: "渠道新客占比\n(=渠道新客/购买人数)",
                            15: "新客成本\n(=媒体花费/渠道新客数)",
                            16: "预估新客GMV\n(=客单价*新客数)",
                            17: "预估整体新客ROI\n(=新客GMV/媒体花费)",
                            18: "全店客单价",
                        }

                    for col, val in header_map.items():
                        ws.cell(row=header_row, column=col).value = val

                    if header_height:
                        ws.row_dimensions[header_row].height = header_height

                    # ===== 步骤 8：写数据行（样式复制原第一数据行） =====
                    for i, (_, row) in enumerate(df_to_write.iterrows()):
                        r = header_row + 1 + i
                        apply_style(r, data_style)

                        ws.cell(row=r, column=2).value = row.get("品牌", "")
                        ws.cell(row=r, column=3).value = "RED KFS" if platform_report == "RED" else "DY"
                        ws.cell(row=r, column=4).value = float(row.get("媒体花费", 0))
                        ws.cell(row=r, column=5).value = float(row.get("阅读数", 0))
                        ws.cell(row=r, column=6).value = ""
                        ws.cell(row=r, column=7).value = float(row.get("进店UV", 0))
                        ws.cell(row=r, column=8).value = ""
                        # 衍生指标 → 公式
                        ws.cell(row=r, column=9).value = f"=G{r}/E{r}"
                        ws.cell(row=r, column=10).value = f"=K{r}/G{r}"
                        ws.cell(row=r, column=11).value = float(row.get("购买人数", 0))
                        ws.cell(row=r, column=12).value = float(row.get("渠道新客数", 0))
                        ws.cell(row=r, column=13).value = ""
                        ws.cell(row=r, column=14).value = f"=L{r}/K{r}"
                        ws.cell(row=r, column=15).value = f"=D{r}/L{r}"
                        ws.cell(row=r, column=16).value = f"=L{r}*R{r}"
                        ws.cell(row=r, column=17).value = f"=P{r}/D{r}"
                        ws.cell(row=r, column=18).value = float(row.get("全店客单价", 0))

                        if data_height:
                            ws.row_dimensions[r].height = data_height
                    # ===== 步骤 9：计算并写入品牌组合 Ave. 行 =====
                    def calc_portfolio(df, brands):
                        sub = df[df["品牌"].isin(brands)]
                        if len(sub) == 0:
                            return None
                        spend = sub["媒体花费"].sum()
                        read = sub["阅读数"].sum()
                        visit = sub["进店UV"].sum()
                        buyer = sub["购买人数"].sum()
                        new_cust = sub["渠道新客数"].sum()
                        gmv = sub["预估新客GMV"].sum()
                        return {
                            "媒体花费": spend, "阅读数": read, "进店UV": visit,
                            "进店率": visit / read if read else 0,
                            "购买转化率": buyer / visit if visit else 0,
                            "购买人数": buyer, "渠道新客数": new_cust,
                            "渠道新客占比": new_cust / buyer if buyer else 0,
                            "新客成本": spend / new_cust if new_cust else 0,
                            "预估新客GMV": gmv,
                            "预估整体新客ROI": gmv / spend if spend else 0,
                            "全店客单价": gmv / new_cust if new_cust else 0,
                        }

                    cat = df_to_write.get("Category", pd.Series([""] * len(df_to_write))).astype(str).str.lower()
                    makeup_brands = df_to_write[cat.str.contains("make", na=False)]["品牌"].tolist()
                    frag_brands = df_to_write[cat.str.contains("fragrance", na=False)]["品牌"].tolist()
                    skincare_brands = df_to_write[cat.str.contains("skincare", na=False)]["品牌"].tolist()
                    all_brands = df_to_write["品牌"].unique().tolist()

                    portfolios = [
                        ("品牌组合 护肤 Ave.", skincare_brands),
                        ("品牌组合 FRAG. Ave.", frag_brands),
                        ("品牌组合 彩妆 Ave.", makeup_brands),
                        ("品牌组合 Ave.", all_brands),
                    ]

                    for i, (label, brands) in enumerate(portfolios):
                        r = header_row + 1 + n_data + i
                        apply_style(r, data_style)
                        ws.cell(row=r, column=7).value = label   # ← 标签写在 G 列（进店UV）
                        res = calc_portfolio(df_to_write, brands)
                        if res:
                            ws.cell(row=r, column=4).value = res["媒体花费"]
                            ws.cell(row=r, column=5).value = res["阅读数"]
                            ws.cell(row=r, column=9).value = res["进店率"]
                            ws.cell(row=r, column=10).value = res["购买转化率"]
                            ws.cell(row=r, column=11).value = res["购买人数"]
                            ws.cell(row=r, column=12).value = res["渠道新客数"]
                            ws.cell(row=r, column=14).value = res["渠道新客占比"]
                            ws.cell(row=r, column=15).value = res["新客成本"]
                            ws.cell(row=r, column=16).value = res["预估新客GMV"]
                            ws.cell(row=r, column=17).value = res["预估整体新客ROI"]
                            ws.cell(row=r, column=18).value = res["全店客单价"]

                    # ===== 步骤 10：所有 Brand 表头行行高固定 45 =====
                    for row in range(1, ws.max_row + 1):
                        if ws.cell(row=row, column=2).value == "Brand":
                            ws.row_dimensions[row].height = 45
                            
                    # ===== 保存 =====
                    output_path = f"output_{platform_report}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
                    wb.save(output_path)

                    with open(output_path, "rb") as f:
                        st.download_button(
                            label="⬇ 下载报告 Excel",
                            data=f,
                            file_name=f"E2E_{platform_report}_报告.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    st.success(f"报告已生成：{n_data} 条 {platform_report} 数据，共插入 {n_rows} 行。")

            except Exception as e:
                st.error(f"生成失败：{str(e)}")
                import traceback
                st.code(traceback.format_exc())

        st.markdown('</div>', unsafe_allow_html=True)
        
# ============ 页面：历史趋势 ============
elif page == "📈 历史趋势":
    st.markdown('<div class="page-title">历史数据与趋势</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">周期：{period_filter} · 平台：{platform_filter} · 品牌：{brand_filter}</div>', unsafe_allow_html=True)
    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### 新客成本趋势")
    st.info("待接入历史数据后，这里显示趋势图。")
    st.markdown('</div>', unsafe_allow_html=True)
    st.markdown('<div class="e2e-card">', unsafe_allow_html=True)
    st.markdown("#### 自动结论")
    st.info("待接入历史数据后，这里自动生成周期结论。")
    st.markdown('</div>', unsafe_allow_html=True)
